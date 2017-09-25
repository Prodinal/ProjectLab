import evaluate_image_slice
import cv2
import plotly as py
import plotly.graph_objs as pgo
import numpy as np
import openpyxl

def Main():
    evaluate_image_slice.init_tf("D:\\tmp\\output_graph.pb", ["positive", "negative"])
    
    test_image = cv2.imread('d:\\School\\2017Onlab1\\Code\\Creating_Training_Images\\EveryImage\\Nodules\\positive\\JPCLN001Nodule298.jpg')
    result = evaluate_image_slice.featureMapOfSlice(test_image)
    
    print(result[0])
    table = [result[0], result[0], result[0]]
    # tryingPlotly([result[0], result[0], result[0]])

    wb = openpyxl.load_workbook('d:\\School\\2017Onlab1\\Code\\GoogleNet\\EvaluateImage\\' + "FeatureMaps.xlsx")
    ws = wb.active
    ws.append(table[0].tolist())
    ws.append(table[1].tolist())
    ws.append(table[2].tolist())

    chart = createChart(ws, mincol=1, maxcol=len(result[0]), minrow=1, maxrow=3)
    #ws.add_chart(chart, 'A' + str(ws.max_row + 1))
    ws.add_chart(chart, 'A76')

    wb.save('d:\\School\\2017Onlab1\\Code\\GoogleNet\\EvaluateImage\\' + "FeatureMapsDebug.xlsx")


def createChart(ws, mincol, maxcol, minrow, maxrow):
    c = openpyxl.chart.LineChart()
    c.title = 'Parallel Coordinates'
    c.style = 13
    c.y_axis.title = 'Value'
    c.x_axis.title = 'Coord'


    dataArea = openpyxl.chart.Reference(ws, min_col = mincol, max_col = maxcol, min_row = minrow, max_row = maxrow)
    c.add_data(dataArea, titles_from_data = True)

    for i in range(maxrow):
        if i <= 1:
            continue
        cell = ws._get_cell(i,1).value
        s = c.series[i-1]
        if cell == 'positive':
            s.marker.graphicalProperties.solidFill = 'FFFF00'
            s.marker.graphicalProperties.line.solidFill = 'FFFF00'
        elif cell == 'negative':
            s.marker.graphicalProperties.solidFill = '0000FF'
            s.marker.graphicalProperties.line.solidFill = '0000FF'
    return c

# tried to use plotly to draw parallel coordinate plots, does not work withouth registration x.x
def tryingPlotly(data):
    maxValues = np.amax(data, axis = 0)
    
    line = dict(color = 'blue')

    columns = []

    for i in range(len(data[0])):
        column = dict(range = [0, maxValues[i]],
                      label = str(i),
                      values = [row[i] for row in data])
        columns.append(column)
    
    drawData = [
        pgo.Parcoords(
            line = line,
            dimensions = list(columns)
        )
    ]

    py.plotly.iplot(drawData, fileName = 'debugParallelCoords')


if __name__ == '__main__':
    Main()